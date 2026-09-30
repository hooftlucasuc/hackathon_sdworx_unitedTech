"""File → text. md/txt/pdf/docx/eml/json. JSON is a Teams/Slack export: one segment per message."""

from __future__ import annotations

import json
import logging
from dataclasses import dataclass
from email import policy
from email.parser import BytesParser
from pathlib import Path
from typing import Optional

log = logging.getLogger(__name__)

SUPPORTED = (".md", ".txt", ".pdf", ".docx", ".eml", ".json")


@dataclass
class Segment:
    """A piece of text with optional chat metadata. Non-chat files yield exactly one segment."""

    text: str
    author: Optional[str] = None
    timestamp: Optional[str] = None


def load(path: Path) -> list[Segment]:
    ext = path.suffix.lower()
    if ext not in SUPPORTED:
        raise ValueError(f"unsupported file type '{ext}'; supported: {', '.join(SUPPORTED)}")
    data = path.read_bytes()
    return load_bytes(data, ext)


def load_bytes(data: bytes, ext: str) -> list[Segment]:
    ext = ext.lower()
    if ext in (".md", ".txt"):
        return [Segment(text=data.decode("utf-8", errors="replace"))]
    if ext == ".pdf":
        return [Segment(text=_pdf(data))]
    if ext == ".docx":
        return [Segment(text=_docx(data))]
    if ext == ".eml":
        return [_eml(data)]
    if ext == ".json":
        return _chat_export(data)
    raise ValueError(f"unsupported file type '{ext}'")


def _pdf(data: bytes) -> str:
    import fitz  # pymupdf

    with fitz.open(stream=data, filetype="pdf") as doc:
        return "\n\n".join(page.get_text() for page in doc)


def _docx(data: bytes) -> str:
    import io

    import docx

    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)


def _eml(data: bytes) -> Segment:
    msg = BytesParser(policy=policy.default).parsebytes(data)
    body = msg.get_body(preferencelist=("plain", "html"))
    text = body.get_content() if body else ""
    header = f"Subject: {msg.get('subject', '')}\nFrom: {msg.get('from', '')}\nDate: {msg.get('date', '')}\n\n"
    return Segment(text=header + text, author=str(msg.get("from", "")) or None, timestamp=str(msg.get("date", "")) or None)


def _chat_export(data: bytes) -> list[Segment]:
    raw = json.loads(data.decode("utf-8", errors="replace"))
    messages = raw.get("messages", raw) if isinstance(raw, dict) else raw
    if not isinstance(messages, list):
        raise ValueError("chat export must be a list of {author, timestamp, text}")
    segments: list[Segment] = []
    for m in messages:
        if not isinstance(m, dict):
            continue
        text = str(m.get("text") or m.get("content") or "").strip()
        if not text:
            continue
        author = m.get("author") or m.get("from") or m.get("user")
        ts = m.get("timestamp") or m.get("date") or m.get("ts")
        prefix = f"[{ts}] {author}: " if (ts or author) else ""
        segments.append(Segment(text=prefix + text, author=str(author) if author else None, timestamp=str(ts) if ts else None))
    if not segments:
        log.warning("chat export contained no messages")
    return segments
